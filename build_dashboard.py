import os
import sys
from pathlib import Path

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent

# Internal preprocessing file used to build the final client workbook.
# This is not part of the client deliverable; dashboard.xlsx is the only Excel file that should be shared externally.
CLEAN_XLSX = BASE_DIR / 'clean_data.xlsx'
OUTPUT_XLSX = BASE_DIR / 'dashboard.xlsx'
AD_SPEND_CSV = BASE_DIR / 'ad_spend.csv'


def auto_fit_column_widths(ws):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col_cells)
        ws.column_dimensions[col_letter].width = min(max(12, max_len + 2), 40)


def find_column(df, keywords):
    lowercase = {str(c).lower(): c for c in df.columns}
    for kw in keywords:
        for lc, actual in lowercase.items():
            if kw.lower() in lc:
                return actual
    return None


def prepare_numeric(df, colname):
    if colname and colname in df.columns:
        blanks = df[colname].isna() | df[colname].astype(str).str.strip().eq('')
        if blanks.any():
            print(f'Warning: {colname} has {int(blanks.sum())} blank/missing values.')
        df[colname] = pd.to_numeric(df[colname], errors='coerce')


def standardize_categorical_columns(df):
    changes = {}

    pay_col = find_column(df, ['payment', 'payment_method', 'payment method'])
    if pay_col and pay_col in df.columns:
        before = df[pay_col].copy()
        s = df[pay_col]
        mask = s.notna()
        normalized = s[mask].astype(str).str.replace(r'[\.\_\s]', '', regex=True).str.upper()
        payment_map = {
            'PAYPAL': 'PayPal',
            'UPI': 'UPI',
            'COD': 'Cash on Delivery',
            'NETBANKING': 'Net Banking',
            'CREDITCARD': 'Credit Card',
            'CREDIT_CARD': 'Credit Card',
            'DEBITCARD': 'Debit Card',
        }
        mapped = normalized.map(payment_map)
        fallback = s[mask].astype(str).str.strip().str.title()
        final = mapped.where(mapped.notna(), fallback)
        df.loc[mask, pay_col] = final
        changes['Payment_Method'] = int((before.fillna('__NA__').astype(str).str.strip() != df[pay_col].fillna('__NA__').astype(str).str.strip()).sum())

    country_col = find_column(df, ['country'])
    if country_col and country_col in df.columns:
        before = df[country_col].copy()
        s = df[country_col]
        mask = s.notna()
        normalized = s[mask].astype(str).str.replace(r'[\.\_\s]', '', regex=True).str.upper()
        country_map = {
            'UAE': 'UAE',
            'USA': 'USA',
            'UNITEDSTATES': 'USA',
            'UNITEDSTATESOFAMERICA': 'USA',
            'IN': 'India',
            'INDIA': 'India',
            'DE': 'Germany',
            'GERMANY': 'Germany',
            'UK': 'UK',
            'UNITEDKINGDOM': 'UK',
        }
        mapped = normalized.map(country_map)
        fallback = s[mask].astype(str).str.strip().str.title()
        final = mapped.where(mapped.notna(), fallback)
        df.loc[mask, country_col] = final
        changes['Country'] = int((before.fillna('__NA__').astype(str).str.strip() != df[country_col].fillna('__NA__').astype(str).str.strip()).sum())

    cat_col = find_column(df, ['category', 'product_category', 'product category'])
    if cat_col and cat_col in df.columns:
        before = df[cat_col].copy()
        s = df[cat_col]
        mask = s.notna()
        df.loc[mask, cat_col] = s[mask].astype(str).str.strip().str.title()
        changes['Product_Category'] = int((before.fillna('__NA__').astype(str).str.strip() != df[cat_col].fillna('__NA__').astype(str).str.strip()).sum())

    for key, value in changes.items():
        print(f'Standardized {key}: {value} rows modified')


def add_box(ws, start_row, start_col, label, value, fill_color='D9EAF7'):
    end_col = start_col + 2
    top_range = f'{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{start_row}'
    bottom_range = f'{get_column_letter(start_col)}{start_row + 1}:{get_column_letter(end_col)}{start_row + 1}'
    ws.merge_cells(top_range)
    top = ws.cell(row=start_row, column=start_col)
    top.value = label
    top.font = Font(bold=True, size=12)
    top.alignment = Alignment(horizontal='center', vertical='center')
    top.fill = PatternFill('solid', fgColor=fill_color)
    top.border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF'),
    )

    ws.merge_cells(bottom_range)
    val_cell = ws.cell(row=start_row + 1, column=start_col)
    val_cell.value = value
    val_cell.font = Font(bold=True, size=15)
    val_cell.alignment = Alignment(horizontal='center', vertical='center')
    val_cell.fill = PatternFill('solid', fgColor='FFFFFF')
    val_cell.border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF'),
    )


def create_chart_table(ws, data, title_name, sheet_name):
    ws_hidden = ws.parent.create_sheet(sheet_name)
    ws_hidden.sheet_state = 'hidden'
    ws_hidden.append(list(data.columns))
    for row in data.itertuples(index=False, name=None):
        ws_hidden.append(row)
    return ws_hidden


def load_ad_spend_data():
    if not AD_SPEND_CSV.exists():
        print('Ad spend file not found:', AD_SPEND_CSV)
        return None

    ad_df = pd.read_csv(AD_SPEND_CSV)
    if ad_df.empty:
        print('Ad spend file is empty:', AD_SPEND_CSV)
        return None

    if 'Date' in ad_df.columns:
        ad_df['Date'] = pd.to_datetime(ad_df['Date'], errors='coerce').dt.strftime('%Y-%m-%d')
    if 'Product_Category' in ad_df.columns:
        ad_df['Product_Category'] = ad_df['Product_Category'].astype(str).str.strip().str.title()
    if 'Ad_Spend' in ad_df.columns:
        ad_df['Ad_Spend'] = pd.to_numeric(ad_df['Ad_Spend'], errors='coerce').fillna(0)

    ad_df = ad_df.dropna(subset=['Date', 'Product_Category']).copy()
    ad_df = ad_df.groupby(['Date', 'Product_Category'], as_index=False)['Ad_Spend'].sum()
    return ad_df


def main():
    if not CLEAN_XLSX.exists():
        print('Cleaned Excel not found:', CLEAN_XLSX)
        sys.exit(1)

    df = pd.read_excel(CLEAN_XLSX, sheet_name='Cleaned Data')

    # date cleaning
    date_col = find_column(df, ['order_date', 'date'])
    if date_col:
        df[date_col] = pd.to_datetime(df[date_col], errors='coerce').dt.strftime('%Y-%m-%d')

    # standardize categories
    standardize_categorical_columns(df)
    pm_col = find_column(df, ['payment', 'payment_method', 'payment method'])
    country_col = find_column(df, ['country'])
    cat_col = find_column(df, ['category', 'product_category', 'product category'])
    print('Unique Payment_Method values after standardization:', sorted(df[pm_col].dropna().unique().tolist()) if pm_col else 'N/A')
    print('Unique Country values after standardization:', sorted(df[country_col].dropna().unique().tolist()) if country_col else 'N/A')
    print('Unique Product_Category values after standardization:', sorted(df[cat_col].dropna().unique().tolist()) if cat_col else 'N/A')

    qty_col = find_column(df, ['quantity', 'qty'])
    price_col = find_column(df, ['unit_price', 'price', 'unit price', 'unit_price_usd'])
    discount_col = find_column(df, ['discount'])
    category_col = find_column(df, ['category'])
    country_col2 = find_column(df, ['country'])
    payment_col = find_column(df, ['payment'])
    status_col = find_column(df, ['status'])

    prepare_numeric(df, qty_col)
    prepare_numeric(df, price_col)
    prepare_numeric(df, discount_col)

    if status_col and status_col in df.columns:
        status_clean = df[status_col].fillna('').astype(str).str.strip().str.title()
        df['Revenue_Type'] = np.where(status_clean.isin(['Delivered', 'Shipped', 'Pending']), 'Fulfilled', 'Cancelled/Returned')
    else:
        df['Revenue_Type'] = 'Fulfilled'

    df['Revenue'] = (
        df.get(qty_col, pd.Series(0, index=df.index)).astype(float)
        * df.get(price_col, pd.Series(0, index=df.index)).astype(float)
        * (1 - df.get(discount_col, pd.Series(0, index=df.index)).fillna(0).astype(float) / 100.0)
    )

    df = df.dropna(subset=[qty_col, price_col]).copy() if qty_col and price_col else df.copy()

    gross_revenue = float(df['Revenue'].sum())
    fulfilled_revenue = float(df.loc[df['Revenue_Type'] == 'Fulfilled', 'Revenue'].sum())
    cancelled_revenue = float(df.loc[df['Revenue_Type'] == 'Cancelled/Returned', 'Revenue'].sum())

    ad_spend_df = load_ad_spend_data()
    ad_performance = None
    if ad_spend_df is not None and not ad_spend_df.empty:
        ad_revenue = df.loc[df['Revenue_Type'] == 'Fulfilled', [date_col, 'Product_Category', 'Revenue']].copy()
        if date_col:
            ad_revenue = ad_revenue.rename(columns={date_col: 'Date'})
        ad_revenue['Date'] = pd.to_datetime(ad_revenue['Date'], errors='coerce').dt.strftime('%Y-%m-%d')
        ad_revenue['Product_Category'] = ad_revenue['Product_Category'].astype(str).str.strip().str.title()
        ad_revenue = ad_revenue.dropna(subset=['Date', 'Product_Category']).groupby(['Date', 'Product_Category'], as_index=False)['Revenue'].sum().rename(columns={'Revenue': 'Fulfilled_Revenue'})
        ad_performance = ad_revenue.merge(ad_spend_df, on=['Date', 'Product_Category'], how='left')
        ad_performance['Ad_Spend'] = ad_performance['Ad_Spend'].fillna(0)
        ad_performance['ROAS'] = np.where(ad_performance['Ad_Spend'] > 0, ad_performance['Fulfilled_Revenue'] / ad_performance['Ad_Spend'], 0)
        ad_performance['Profit_After_Ads'] = ad_performance['Fulfilled_Revenue'] - ad_performance['Ad_Spend']
        ad_performance = ad_performance[ad_performance['Ad_Spend'] > 0].copy()
        print('Ad performance rows merged:', len(ad_performance))
        print('Average ROAS:', round(float(ad_performance['ROAS'].mean()), 2))

    # active period trimming for trend smoothness
    df_trim = df.copy()
    if date_col and date_col in df_trim.columns:
        df_trim[date_col] = pd.to_datetime(df_trim[date_col], errors='coerce')
        df_trim['YearMonth'] = df_trim[date_col].dt.to_period('M').dt.to_timestamp()
        monthly_counts = df_trim.groupby('YearMonth').size().reset_index(name='count')
        active_months = monthly_counts[monthly_counts['count'] > 50]
        if not active_months.empty:
            last_active = active_months['YearMonth'].max()
            df_trim = df_trim[df_trim['YearMonth'] <= last_active].copy()
            print('Trimmed date range to active period:', df_trim[date_col].min().date(), '-', df_trim[date_col].max().date())

    # build workbook from scratch
    wb = Workbook()
    ws_data = wb.active
    ws_data.title = 'Data'

    df_out = df.copy()
    if date_col:
        df_out.rename(columns={date_col: 'Order_Date'}, inplace=True)
    if category_col:
        df_out.rename(columns={category_col: 'Product_Category'}, inplace=True)
    if country_col2:
        df_out.rename(columns={country_col2: 'Country'}, inplace=True)
    if payment_col:
        df_out.rename(columns={payment_col: 'Payment_Method'}, inplace=True)
    if status_col:
        df_out.rename(columns={status_col: 'Order_Status'}, inplace=True)

    # keep only relevant columns for data filtering
    keep_cols = ['Order_ID', 'Customer_ID', 'Order_Date', 'Product_Category', 'Product_Name', 'Quantity', 'Unit_Price_USD', 'Discount_Percent', 'Payment_Method', 'Shipping_City', 'Country', 'Order_Status', 'Customer_Rating', 'Revenue_Type', 'Revenue']
    for c in keep_cols:
        if c not in df_out.columns:
            df_out[c] = np.nan
    df_out = df_out[keep_cols]

    # write headers and rows to Data sheet
    ws_data.append(list(df_out.columns))
    for row in df_out.itertuples(index=False, name=None):
        ws_data.append(row)

    # add official Excel Table with AutoFilters
    table = Table(displayName='SalesData', ref=f'A1:{get_column_letter(ws_data.max_column)}{ws_data.max_row}')
    table.tableStyleInfo = TableStyleInfo(
        name='TableStyleMedium9',
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws_data.add_table(table)
    auto_fit_column_widths(ws_data)
    ws_data.freeze_panes = 'A2'

    # dashboard sheet
    ws_dash = wb.create_sheet('Dashboard')
    ws_dash.sheet_view.showGridLines = False
    ws_dash.merge_cells('A1:F1')
    ws_dash['A1'] = 'Sales Dashboard'
    ws_dash['A1'].font = Font(bold=True, size=18, color='FFFFFF')
    ws_dash['A1'].fill = PatternFill('solid', fgColor='1F4E78')
    ws_dash['A1'].alignment = Alignment(horizontal='center', vertical='center')

    # KPI boxes
    kpis = [('Gross Revenue', gross_revenue), ('Fulfilled Revenue', fulfilled_revenue), ('Cancelled Revenue', cancelled_revenue)]
    for idx, (label, value) in enumerate(kpis):
        start_row = 3 + idx * 5
        add_box(ws_dash, start_row, 1, label, f'${value:,.2f}')

    if ad_performance is not None and not ad_performance.empty:
        total_ad_spend = float(ad_performance['Ad_Spend'].sum())
        avg_roas = float(ad_performance['ROAS'].mean())
        profit_after_ads = float(ad_performance['Profit_After_Ads'].sum())

        ws_dash.merge_cells('H1:K1')
        ws_dash['H1'] = 'Ad Performance'
        ws_dash['H1'].font = Font(bold=True, size=16, color='FFFFFF')
        ws_dash['H1'].fill = PatternFill('solid', fgColor='2F75B5')
        ws_dash['H1'].alignment = Alignment(horizontal='center', vertical='center')

        ad_kpis = [('Ad Spend', total_ad_spend), ('Avg ROAS', avg_roas), ('Profit After Ads', profit_after_ads)]
        for idx, (label, value) in enumerate(ad_kpis):
            start_row = 3 + idx * 5
            formatted_value = f'${value:,.2f}' if label != 'Avg ROAS' else f'{value:,.2f}x'
            add_box(ws_dash, start_row, 8, label, formatted_value)

        ad_detail = ad_performance[['Date', 'Product_Category', 'Fulfilled_Revenue', 'Ad_Spend', 'ROAS', 'Profit_After_Ads']].copy()
        ad_detail_ws = create_chart_table(ws_dash, ad_detail, 'Ad Detail', 'Ad_Performance_Data')

        ad_chart_data = ad_performance.groupby('Product_Category', as_index=False)['ROAS'].mean().sort_values('ROAS', ascending=False)
        ad_chart_ws = create_chart_table(ws_dash, ad_chart_data, 'Ad Summary', 'Ad_Performance_Summary')

        ad_chart = BarChart()
        ad_chart.type = 'col'
        ad_chart.style = 10
        ad_chart.title = 'Avg ROAS by Category'
        ad_chart.y_axis.title = 'ROAS'
        ad_chart.x_axis.title = 'Category'
        ad_chart.height = 14
        ad_chart.width = 22
        ad_chart.y_axis.number_format = '0.00x'
        ad_data_ref = Reference(ad_chart_ws, min_col=2, min_row=1, max_row=ad_chart_ws.max_row)
        ad_cats = Reference(ad_chart_ws, min_col=1, min_row=2, max_row=ad_chart_ws.max_row)
        ad_chart.add_data(ad_data_ref, titles_from_data=True)
        ad_chart.set_categories(ad_cats)
        ws_dash.add_chart(ad_chart, 'B90')

    # helper data for charts
    cat_data = df_trim.groupby('Product_Category' if 'Product_Category' in df_trim.columns else category_col, dropna=False)['Revenue'].sum().reset_index().sort_values('Revenue', ascending=False)
    cat_data.columns = ['Product_Category', 'Revenue']
    cat_ws = create_chart_table(ws_dash, cat_data, 'Category', 'Category_Data')

    country_data = df_trim.groupby('Country' if 'Country' in df_trim.columns else country_col2, dropna=False)['Revenue'].sum().reset_index().sort_values('Revenue', ascending=False)
    country_data.columns = ['Country', 'Revenue']
    country_ws = create_chart_table(ws_dash, country_data, 'Country', 'Country_Data')

    payment_data = df_trim.groupby('Payment_Method' if 'Payment_Method' in df_trim.columns else payment_col, dropna=False)['Revenue'].sum().reset_index().sort_values('Revenue', ascending=False)
    payment_data.columns = ['Payment_Method', 'Revenue']
    payment_ws = create_chart_table(ws_dash, payment_data, 'Payment', 'Payment_Data')

    trend_data = df_trim.copy()
    if date_col and date_col in trend_data.columns:
        trend_data['Month'] = pd.to_datetime(trend_data[date_col], errors='coerce').dt.to_period('M').dt.to_timestamp()
        trend_data = trend_data.groupby('Month', as_index=False)['Revenue'].sum()
    else:
        trend_data = pd.DataFrame({'Month': pd.to_datetime(df_trim.index), 'Revenue': df_trim['Revenue'].values})
    trend_ws = create_chart_table(ws_dash, trend_data, 'Trend', 'Trend_Data')

    # category chart
    category_chart = BarChart()
    category_chart.type = 'col'
    category_chart.style = 10
    category_chart.title = 'Revenue by Category'
    category_chart.y_axis.title = 'Revenue'
    category_chart.x_axis.title = 'Category'
    category_chart.height = 14
    category_chart.width = 22
    category_chart.y_axis.number_format = '$#,##0'
    category_data_ref = Reference(cat_ws, min_col=2, min_row=1, max_row=cat_ws.max_row)
    category_cats = Reference(cat_ws, min_col=1, min_row=2, max_row=cat_ws.max_row)
    category_chart.add_data(category_data_ref, titles_from_data=True)
    category_chart.set_categories(category_cats)
    ws_dash.add_chart(category_chart, 'B10')

    # country chart
    country_chart = BarChart()
    country_chart.type = 'bar'
    country_chart.style = 10
    country_chart.title = 'Revenue by Country'
    country_chart.y_axis.title = 'Revenue'
    country_chart.x_axis.title = 'Country'
    country_chart.height = 14
    country_chart.width = 22
    country_chart.y_axis.number_format = '$#,##0'
    country_data_ref = Reference(country_ws, min_col=2, min_row=1, max_row=country_ws.max_row)
    country_cats = Reference(country_ws, min_col=1, min_row=2, max_row=country_ws.max_row)
    country_chart.add_data(country_data_ref, titles_from_data=True)
    country_chart.set_categories(country_cats)
    ws_dash.add_chart(country_chart, 'B30')

    # payment chart
    payment_chart = DoughnutChart()
    payment_chart.title = 'Revenue by Payment Method'
    payment_chart.style = 10
    payment_chart.height = 14
    payment_chart.width = 22
    payment_data_ref = Reference(payment_ws, min_col=2, min_row=1, max_row=payment_ws.max_row)
    payment_cats = Reference(payment_ws, min_col=1, min_row=2, max_row=payment_ws.max_row)
    payment_chart.add_data(payment_data_ref, titles_from_data=True)
    payment_chart.set_categories(payment_cats)
    ws_dash.add_chart(payment_chart, 'B50')

    # monthly trend chart
    trend_chart = LineChart()
    trend_chart.title = 'Monthly Revenue Trend'
    trend_chart.style = 13
    trend_chart.height = 14
    trend_chart.width = 22
    trend_chart.y_axis.title = 'Revenue'
    trend_chart.x_axis.title = 'Month'
    trend_chart.y_axis.number_format = '$#,##0'
    trend_chart.x_axis.number_format = 'mmm yyyy'
    trend_data_ref = Reference(trend_ws, min_col=2, min_row=1, max_row=trend_ws.max_row)
    trend_cats = Reference(trend_ws, min_col=1, min_row=2, max_row=trend_ws.max_row)
    trend_chart.add_data(trend_data_ref, titles_from_data=True)
    trend_chart.set_categories(trend_cats)
    ws_dash.add_chart(trend_chart, 'B70')

    # style dashboard header/elements
    for row in ws_dash.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = Font(bold=True, color='FFFFFF')

    wb.save(OUTPUT_XLSX)
    print(f'Wrote fully automated dashboard to {OUTPUT_XLSX}')


if __name__ == '__main__':
    main()
